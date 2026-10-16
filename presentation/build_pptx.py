"""Build the project presentation on the NMIMS template (docs/templates/Presentation_Format.pptx).

    python presentation/build_pptx.py      # -> presentation/OEE_Smart_Factory_Presentation.pptx

Structure follows the SRS (docs/SRS.docx / the submitted SRS): requirements -> design (UML, DFD) -> implementation -> verification.
All diagrams are native PowerPoint shapes (editable, always sharp); numbers come from the measured test runs.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from lxml import etree
from pptx import Presentation

import deckkit as K
from deckkit import *  # noqa: F401,F403

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "docs" / "templates" / "Presentation_Format.pptx"
OUT = ROOT / "presentation" / "OEE_Smart_Factory_Presentation.pptx"
CROP = ROOT / "presentation" / "_crops"
CROP.mkdir(exist_ok=True)

PROJECT_TITLE = "Real-Time OEE and Anomaly Dashboard for a Smart Factory System"

prs = Presentation(str(TEMPLATE))
LAYOUT = prs.slide_layouts.get_by_name("Title and Content")


class Ctx:
    """Shared state handed to the slide modules."""
    prs = prs
    root = ROOT
    crop_dir = CROP


def _strip_placeholders(slide):
    for sh in list(slide.placeholders):
        t = str(sh.placeholder_format.type)
        if "SLIDE_NUMBER" not in t:
            sh._element.getparent().remove(sh._element)


def slide_number(s):
    """Right-aligned page number at exactly the template's own position."""
    box = s.shapes.add_textbox(E(8.83), E(6.97), E(2.33), E(0.4))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"))
    fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
    fld.set("type", "slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr"))
    rpr.set("lang", "en-US"), rpr.set("sz", "1200")
    sf = etree.SubElement(rpr, qn("a:solidFill"))
    clr = etree.SubElement(sf, qn("a:srgbClr"))
    clr.set("val", "8C8C8C")
    lat = etree.SubElement(rpr, qn("a:latin"))
    lat.set("typeface", FONT)
    t = etree.SubElement(fld, qn("a:t"))
    t.text = "‹#›"


def new_slide(title: str, notes: str):
    s = prs.slides.add_slide(LAYOUT)
    _strip_placeholders(s)
    K.set_current(title)
    tb(s, X0, 0.1, 10.9, 0.55, title, size=28, bold=True, color=WHITE, anchor="m")
    slide_number(s)
    s.notes_slide.notes_text_frame.text = notes
    return s


Ctx.new = staticmethod(new_slide)

import s_intro   # noqa: E402
import s_design  # noqa: E402
import s_result  # noqa: E402

# ============================================================================ slide 1: template title
s1 = prs.slides[0]
for sh in s1.shapes:
    if sh.has_text_frame and "TITLE OF PROJECT" in sh.text_frame.text:
        tf = sh.text_frame
        lines = ["A Presentation on", f"“{PROJECT_TITLE}”", "", "By", "Sanika Jamadar  (Roll No. B291 · 70022400523, Batch A2)", "",
                 "Subject: Software Engineering", "Under the Guidance of  Dr. Naziya Hussain"]
        for i, p in enumerate(list(tf.paragraphs)):
            if i > 0:
                p._p.getparent().remove(p._p)
        first = tf.paragraphs[0]
        for r in list(first.runs):
            r._r.getparent().remove(r._r)
        for i, ln in enumerate(lines):
            p = first if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = ln
            r.font.name, r.font.bold = FONT, True
            r.font.size = Pt(26 if i == 1 else 20)
            r.font.color.rgb = INK
s1.notes_slide.notes_text_frame.text = (
    "Introduce yourself and the project in one sentence: a software-only smart-factory monitor that shows OEE live and warns operators "
    "three different ways - rules, a machine-learning model and a trend projection - before a machine fails. It is built exactly from my SRS.")

# ============================================================================ slide 2: agenda (template slide, rebuilt on the grid)
s2 = prs.slides[1]
for sh in list(s2.shapes):
    if sh.is_placeholder and "SLIDE_NUMBER" not in str(sh.placeholder_format.type):
        sh._element.getparent().remove(sh._element)
K.set_current("Agenda")
tb(s2, X0, 0.1, 10.9, 0.55, "Agenda", size=28, bold=True, color=WHITE, anchor="m")
agenda = [("Introduction & Problem Statement", "slides 3 – 5"), ("Objectives & Existing vs. Proposed", "slides 6 – 7"),
          ("Users & Requirements (SRS)", "slides 8 – 9"), ("System Architecture & Data Flow", "slides 10 – 13"),
          ("Methodology & Modules", "slides 14 – 15"), ("Technologies Used", "slide 16"), ("OEE, Alerts, ML & What-If", "slides 17 – 21"),
          ("UML & Database Design", "slides 22 – 27"), ("Deployment & Run Modes", "slides 28 – 29"), ("Implementation / Screenshots", "slides 30 – 31"),
          ("Live Demo", "slide 32"), ("Results, Testing & Traceability", "slides 33 – 34"), ("Advantages, Limitations & Applications", "slide 35"),
          ("Conclusion, Future Scope & References", "slides 36 – 37")]
colw = (W - 0.4) / 2
rows = 7
rh = (Y1 - 1.25) / rows
for i, (t, r) in enumerate(agenda):
    c, rr = divmod(i, rows)
    x = X0 + c * (colw + 0.4)
    y = 1.25 + rr * rh
    box(s2, x, y, colw, rh - 0.14, "", fill="grey", edge="white", radius=0.1)
    badge(s2, x + 0.38, y + (rh - 0.14) / 2, 0.46, f"{i + 1}", size=15)
    tb(s2, x + 0.8, y + 0.07, colw - 0.9, 0.36, [{"text": t, "size": 17, "bold": True}], anchor="m")
    tb(s2, x + 0.8, y + 0.42, colw - 0.9, 0.2, r, size=12, color=MUTED, anchor="m")
s2.notes_slide.notes_text_frame.text = ("Roadmap in the order of the SRS: requirements first, then design (architecture, UML, data flow), then how it was built, "
                                        "then how it was verified. About 15 minutes plus a 3-minute live demo.")

# ============================================================================ content slides
s_intro.build(Ctx)
s_design.build(Ctx)
s_result.build(Ctx)

# ============================================================================ template "Thank you" slide goes last
sld = prs.slides._sldIdLst
ids = list(sld)
thank = ids[2]
sld.remove(thank)
sld.append(thank)
for sh in prs.slides[len(prs.slides) - 1].shapes:
    if sh.has_text_frame and "Thank" in sh.text_frame.text:
        sh.text_frame.paragraphs[0].runs[0].text = "Thank You !!!  Questions?"
        sh.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
prs.slides[len(prs.slides) - 1].notes_slide.notes_text_frame.text = (
    "Invite questions. Likely ones: Why Isolation Forest? (unsupervised, no fault labels needed, very fast inference). How is the 85 percent calibrated? "
    "(normal tail about 70 percent, median anomaly about 95 percent). What if the broker or ML fails? (automatic reconnect; last score held; red rules keep working). "
    "What was verified on real databases? (PostgreSQL and InfluxDB in Docker: data written, read back and still there after a restart).")

prs.save(str(OUT))
print("saved", OUT, "slides:", len(prs.slides))
if K.WARNINGS:
    print("\nPOSSIBLE TEXT OVERFLOW (check these slides):")
    for w in K.WARNINGS:
        print("  -", w)
