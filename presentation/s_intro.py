"""Slides 3-9: overview, introduction, problem, objectives, existing vs proposed, users, requirements."""
from __future__ import annotations

from pptx.util import Pt

from deckkit import *  # noqa: F401,F403


def build(ctx):
    new = ctx.new

    # ------------------------------------------------------------------ 3  the project in 30 seconds
    s = new("The Project in 30 Seconds",
            "If the faculty remembers one slide, it should be this one. A car dashboard shows speed and warns you before the engine overheats. "
            "Our software does that for a factory: virtual machines send readings, a smart backend computes productivity (OEE) and watches for trouble in three different ways, "
            "and the operator sees a four-colour status and can act on it.")
    tb(s, X0, 1.1, W, 0.8, ["Think of a **car dashboard – for a factory.** It shows how well every machine is running and warns the operator **before** something breaks."], size=20)
    names = [("STEP 1", "Virtual machines", "blue", "Five machines send temperature, vibration, status and cycle count every 0.5 s."),
             ("STEP 2", "MQTT broker", "orange", "A lightweight message broker (Mosquitto) carries every reading."),
             ("STEP 3", "Smart backend", "green", "Computes OEE and checks three detectors: rules, machine learning, trend line."),
             ("STEP 4", "Live dashboard", "cyan", "Shows each machine Green, Yellow, Orange or Red – live, in under a second."),
             ("STEP 5", "Operator", "violet", "Acknowledges, assigns a technician, resolves and logs the maintenance.")]
    gap = 0.45
    nw = (W - 4 * gap) / 5
    for i, (st, nm, col, cap) in enumerate(names):
        x = X0 + i * (nw + gap)
        box(s, x, 2.1, nw, 1.1, [{"text": st, "size": 12, "color": MUTED, "bold": True}, {"text": nm, "size": 16, "bold": True}], fill=col, edge=col)
        tb(s, x, 3.38, nw, 1.5, cap, size=14)
        if i < 4:
            line(s, x + nw + 0.06, 2.65, x + nw + gap - 0.06, 2.65, color=INK, w=2.25)
    tb(s, X0, 5.0, W, 0.35, "What the operator sees on every machine tile", size=17, bold=True)
    chips = [("green", "Green", "Normal operation"), ("yellow", "Yellow", "ML caution: something looks unusual"),
             ("orange", "Orange", "Projected fault: limit reached in ≤ 60 s"), ("red", "Red", "Hard fault: limit crossed")]
    for (x, w), (col, nm, d) in zip(cols(4), chips):
        box(s, x, 5.5, w, 1.2, [{"text": nm, "size": 17, "bold": True}, {"text": d, "size": 14}], fill=col, edge=col, line=2)

    # ------------------------------------------------------------------ 4  introduction: what is OEE
    s = new("Introduction – What is OEE?",
            "OEE - Overall Equipment Effectiveness - is the standard productivity number in manufacturing (ISO 22400). It multiplies three simple questions: "
            "was the machine running, was it fast, and were the parts good? The worked example at the bottom is easy to follow on the board.")
    xs = [(X0 + i * (2.4 + 0.45), 2.4) for i in range(4)]
    items = [("Availability", "run time ÷ planned time", "blue"), ("Performance", "actual speed ÷ ideal speed", "cyan"),
             ("Quality", "good parts ÷ all parts", "violet"), ("OEE", "overall effectiveness", "green")]
    for (x, w), (nm, fm, col) in zip(xs, items):
        box(s, x, 1.2, w, 1.25, [{"text": nm, "size": 22, "bold": True}, {"text": fm, "size": 14}], fill=col, edge=col)
    for i, sym in enumerate(["×", "×", "="]):
        label(s, xs[i][0] + 2.4, 1.2, 0.45, sym, size=30, color=INK, bold=True, h=1.25)
    plain = ["Was the machine running when it should? Stops lower it.",
             "Did it run at full speed? Slow cycles lower it.",
             "How many parts were good? Scrap and rework lower it.",
             "One number from 0 to 100 % – the standard productivity KPI (ISO 22400)."]
    for (x, w), t in zip(xs, plain):
        tb(s, x, 2.62, w, 1.4, t, size=14)
    box(s, X0, 4.25, W, 2.45, "", fill="grey", edge="white")
    tb(s, X0 + 0.3, 4.4, 3.0, 0.4, "Worked example", size=18, bold=True)
    tb(s, X0 + 0.3, 4.85, 2.7, 1.2, "A machine that is up 90 % of the time, runs at 95 % of ideal speed and makes 98 % good parts is not “90 % productive”:", size=14)
    vals = [("90 %", "Availability"), ("×", ""), ("95 %", "Performance"), ("×", ""), ("98 %", "Quality"), ("=", ""), ("83.8 %", "OEE")]
    x = X0 + 3.4
    for v, lab in vals:
        w = 1.35 if lab else 0.55
        if lab:
            big = lab == "OEE"
            tb(s, x, 4.75, w, 0.7, v, size=32, bold=True, color=GREEN_D if big else INK, align="c", anchor="m")
            tb(s, x, 5.5, w, 0.3, lab, size=13, color=MUTED, align="c")
        else:
            tb(s, x, 4.75, w, 0.7, v, size=30, bold=True, color=SLATE, align="c", anchor="m")
        x += w
    tb(s, X0 + 0.3, 6.15, W - 0.6, 0.35, "Reference point often quoted for “world-class” plants: about 85 % (Nakajima, TPM).", size=14, color=MUTED, italic=True)

    # ------------------------------------------------------------------ 5  problem statement
    s = new("Problem Statement",
            "Left: what plants live with today - OEE computed after the shift, alarms only when a limit is already crossed, slow drifts unnoticed. "
            "Right: the four things a manager or engineer actually needs. The green banner is the numeric targets from the SRS that we later prove with measurements.")
    cl, cr = X0, X0 + 5.15
    lw, rw = 4.7, X1 - (X0 + 5.15)
    tb(s, cl, 1.1, lw, 0.4, "Today", size=19, bold=True, color=RED)
    tb(s, cr, 1.1, rw, 0.4, "What is needed", size=19, bold=True, color=GREEN_D)
    pains = ["OEE is worked out **after** the shift, in spreadsheets", "Alarms fire only when a fixed limit is **already crossed**", "Slow drifts and unusual patterns go unnoticed until a breakdown"]
    for i, p in enumerate(pains):
        box(s, cl, 1.6 + i * 1.5, lw, 1.3, p, fill="red", edge="red", size=16, align="l", margin=0.2)
    needs = ["Show **OEE** per machine and for the whole plant, live", "Flag **hard faults** instantly (Red)", "Detect **subtle anomalies** that limits miss (Yellow)",
             "**Predict** a breach early enough to act (Orange)"]
    for i, n in enumerate(needs):
        y = 1.6 + i * 1.125
        box(s, cr, y, rw, 0.97, "", fill="green", edge="green")
        badge(s, cr + 0.45, y + 0.485, 0.5, str(i + 1), fill=GREEN_D)
        tb(s, cr + 0.95, y, rw - 1.1, 0.97, n, size=16, anchor="m")
    right_arrow_big(s, X0 + 4.77, 3.55, 0.33, 0.6)
    box(s, X0, 6.2, W, 0.62, "SRS targets:  updates ≤ 1 s  ·  ML ≤ 200 ms  ·  MQTT reconnect ≤ 5 s  ·  50 clients  ·  5 → 20 machines",
        fill="grey", edge="white", size=15, bold=True, color=GREEN_D)

    # ------------------------------------------------------------------ 6  objectives
    s = new("Objectives",
            "Six objectives; each maps to requirements in the SRS (FR-01 to FR-12) shown in small print, and each is verified by a test - see the traceability slide later.")
    objs = [("Simulate", "5 virtual machines (scalable to 20) send telemetry every 500 ms.", "FR-01 · FR-11"),
            ("Measure", "OEE = Availability × Performance × Quality over rolling 1 h and 8 h windows.", "FR-02"),
            ("Detect", "Red by rules, Yellow by machine learning (> 85 %), Orange by What-If trend.", "FR-03 · FR-04 · FR-07"),
            ("Act", "Acknowledge, assign a technician, resolve, log maintenance, give ML feedback.", "FR-06"),
            ("Report", "7-day OEE trend, downtime per shift and ML accuracy.", "FR-08"),
            ("Run anywhere", "One command to start; validated input, login for writes, documented API.", "FR-09 · FR-10 · FR-12")]
    rowsy = [1.2, 1.2 + 2.6 + GAP]
    for i, (t, d, ids) in enumerate(objs):
        (x, w), y = cols(3)[i % 3], rowsy[i // 3]
        box(s, x, y, w, 2.6, "", fill="grey", edge="white")
        badge(s, x + 0.45, y + 0.45, 0.55, str(i + 1))
        tb(s, x + 0.9, y + 0.2, w - 1.1, 0.5, t, size=20, bold=True, anchor="m")
        tb(s, x + 0.25, y + 0.95, w - 0.5, 1.2, d, size=16)
        tb(s, x + 0.25, y + 2.2, w - 0.5, 0.3, ids, size=12, color=RED, bold=True)

    # ------------------------------------------------------------------ 7  existing vs proposed
    s = new("Existing vs. Proposed System",
            "Contrast: today OEE is retrospective and alarms are fixed thresholds. Our system is live, layered and explainable, and runs with no hardware at all.")
    rows = [("", "Existing practice", "Proposed system"),
            ("OEE", "Spreadsheets after the shift", "Live 1 h / 8 h, with A, P and Q shown separately"),
            ("Fault detection", "Fixed alarm limits only", "Rules + Isolation Forest + What-If projection"),
            ("Warning time", "None – the alarm sounds at the breach", "Orange up to 60 s ahead, with an ETA countdown"),
            ("Operator workflow", "Paper notes / chat", "Acknowledge, assign, resolve, feedback, maintenance log"),
            ("Deployment", "Vendor systems, hardware needed", "Open-source stack, one command, no hardware")]
    shape = s.shapes.add_table(len(rows), 3, E(X0), E(1.2), E(W), E(0.6 + 5 * 0.82))
    tbl = shape.table
    for c, w in enumerate((2.45, 3.9, 4.6)):
        tbl.columns[c].width = E(w)
    for r, row in enumerate(rows):
        tbl.rows[r].height = E(0.6 if r == 0 else 0.82)
        for c, v in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = E(0.18)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = v
            run.font.name, run.font.size, run.font.bold = FONT, Pt(17), r == 0 or c == 0
            run.font.color.rgb = WHITE if r == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = RED if r == 0 else (SOFT["green"] if c == 2 else SOFT["grey"])
    tb(s, X0, 6.25, W, 0.5, "Key idea:  warn **before** the limit is crossed – and make every warning explainable.", size=19, color=GREEN_D)

    # ------------------------------------------------------------------ 8  users & user requirements
    s = new("Users & User Requirements (SRS §3.3, §7.1)",
            "Four user classes from the SRS. Each card shows the user stories (UR-01 to UR-06) that drive the features. The access line shows what each role may do in the app: "
            "everyone can view; operating (acknowledge, maintenance, feedback) needs a login with the right role; control (What-If, simulator) is for the floor manager and administrator.")
    people = [("Floor Manager", "UR-01 · UR-04", ["See all 5 machines with live status and OEE at a glance", "Trigger What-If scenarios for maintenance planning"], "view · operate · control", "blue"),
              ("Maintenance Engineer", "UR-02 · UR-03", ["Receive Red alerts and early Yellow cautions", "Acknowledge, log root cause, mark resolved (audit trail)"], "view · operate", "orange"),
              ("Production Analyst", "UR-05", ["Review OEE trends, downtime by shift and ML accuracy", "Make data-driven decisions on long-term performance"], "view · reports", "violet"),
              ("System Administrator", "UR-06", ["Start the whole stack with one command", "Scale from 5 to 20 simulated machines"], "view · operate · control", "green")]
    for (x, w), (nm, ur, bl, acc, col) in zip(cols(4), people):
        box(s, x, 1.2, w, 4.25, "", fill=col, edge=col)
        tb(s, x + 0.18, 1.35, w - 0.36, 0.7, nm, size=19, bold=True)
        tb(s, x + 0.18, 2.05, w - 0.36, 0.3, ur, size=13, bold=True, color=RED)
        tb(s, x + 0.18, 2.5, w - 0.36, 2.0, [{"text": b, "size": 14, "bullet": True, "space_after": 8} for b in bl], size=14)
        tb(s, x + 0.18, 4.6, w - 0.36, 0.25, "ACCESS IN THE APP", size=11, bold=True, color=MUTED)
        tb(s, x + 0.18, 4.9, w - 0.36, 0.4, acc, size=14, bold=True)
    box(s, X0, 5.7, W, 1.05, [{"text": "Also in the SRS", "size": 15, "bold": True},
                              {"text": "Project supervisors and evaluators (verify the requirements)  ·  Time-Server (system actor for time-window processing)", "size": 14}],
        fill="grey", edge="white", align="l", margin=0.25)

    # ------------------------------------------------------------------ 9  functional requirements
    s = new("Functional Requirements at a Glance (SRS §4)",
            "The twelve functional requirements FR-01 to FR-12 from the SRS on one slide. The strip at the bottom is the non-functional targets. "
            "Later slides prove each one with a test or a measurement.")
    frs = [("FR-01", "Virtual Factory", "5 machines · 500 ms · MQTT over WebSockets"), ("FR-02", "OEE Calculation", "A × P × Q over rolling 1 h and 8 h windows"),
           ("FR-03", "Rule-Based Alerting", "Hard limits, e.g. temperature > 85 °C → Red"), ("FR-04", "ML Anomaly Scoring", "Isolation Forest score 0–100 %; > 85 % → Yellow"),
           ("FR-05", "Interactive Dashboard", "React 2D view · Green / Yellow / Orange / Red"), ("FR-06", "Maintenance Logging", "Acknowledge, assign, notes, ML feedback (CRUD)"),
           ("FR-07", "What-If Simulator", "3-min temperature rise · y = mx + c · Orange 60 s ahead"), ("FR-08", "Historical Reporting", "7-day OEE · downtime per shift · ML accuracy"),
           ("FR-09", "Telemetry Validation", "Rejects negative temperature or invalid status"), ("FR-10", "Authentication", "Login required for all write operations"),
           ("FR-11", "Scalability", "5 → 20 machines by configuration only"), ("FR-12", "Auto-Documentation", "OpenAPI / Swagger at /docs")]
    ch = 1.42
    for i, (fid, nm, d) in enumerate(frs):
        (x, w), y = cols(4)[i % 4], 1.2 + (i // 4) * (ch + 0.2)
        box(s, x, y, w, ch, "", fill="blue", edge="blue")
        tb(s, x + 0.15, y + 0.12, w - 0.3, 0.3, fid, size=13, bold=True, color=RED)
        tb(s, x + 0.15, y + 0.42, w - 0.3, 0.3, nm, size=15, bold=True)
        tb(s, x + 0.15, y + 0.78, w - 0.3, 0.6, d, size=13)
    box(s, X0, 6.0, W, 0.8, [{"text": "Non-functional targets:  updates ≤ 1 s  ·  ML inference ≤ 200 ms  ·  charts ≤ 2 s  ·  reconnect ≤ 5 s  ·  50 clients  ·  5 → 20 machines  ·  "
                                       "retention 30 days raw / 1 year aggregated", "size": 14, "bold": True}], fill="grey", edge="white", color=GREEN_D)
