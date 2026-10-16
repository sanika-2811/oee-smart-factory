"""Slides 10-29: architecture, data flow, DFDs, method, modules, stack, OEE/alerts/ML/What-If, UML, database, deployment."""
from __future__ import annotations

import math

from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_AXIS_CROSSES, XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.oxml.ns import qn
from pptx.util import Pt

from deckkit import *  # noqa: F401,F403


def _white_label(s, cx, cy, w, text, size=12, h=0.26):
    """Label with a white plate so it stays readable on top of a line."""
    b = box(s, cx - w / 2, cy - h / 2, w, h, text, fill="white", edge=None, size=size, color=MUTED, margin=0.03, radius=0.04)
    b.line.fill.background()
    return b


def _ellipse_pt(cx, cy, rx, ry, tx, ty):
    dx, dy = tx - cx, ty - cy
    k = 1.0 / math.sqrt((dx / rx) ** 2 + (dy / ry) ** 2)
    return cx + dx * k, cy + dy * k


def build(ctx):
    new = ctx.new

    # ================================================================== 10  architecture
    s = new("System Architecture",
            "Walk left to right: the simulator publishes JSON over MQTT (WebSockets) to Mosquitto; the FastAPI backend validates, computes OEE and runs the three detectors, "
            "stores data and pushes a snapshot every 500 ms to the React dashboard. Two switches in .env choose the ingest path (direct or mqtt) and the storage (memory or db).")
    bx = X0 + 4.8
    box(s, bx, 1.1, 3.75, 4.1, "", fill="grey", edge="blue", dash=True, line=2)
    tb(s, bx + 0.15, 1.18, 3.45, 0.35, "FastAPI backend (asyncio)", size=16, bold=True, color=EDGE["blue"])
    mods = [("Ingest + validation", "FR-09", "blue"), ("OEE engine", "FR-02 · A × P × Q", "green"), ("Rule engine → RED", "FR-03", "red"),
            ("Isolation Forest → YELLOW", "FR-04", "yellow"), ("What-If trend → ORANGE", "FR-07", "orange"), ("Alerts, maintenance, API", "FR-05 · FR-06", "violet")]
    mw = (3.75 - 0.2 - 0.15) / 2
    for i, (t, fr, col) in enumerate(mods):
        mx = bx + 0.1 + (i % 2) * (mw + 0.15)
        my = 1.65 + (i // 2) * 1.17
        box(s, mx, my, mw, 1.05, [{"text": t, "size": 13, "bold": True}, {"text": fr, "size": 12, "color": MUTED}], fill=col, edge=col, margin=0.08)
    side = [("Virtual Factory Simulator", "5–20 machines · every 500 ms", X0, 1.55, "blue"), ("Eclipse Mosquitto", "MQTT over WebSockets", X0 + 2.4, 1.55, "orange"),
            ("React 18 dashboard", "Chart.js 2D charts", X0 + 9.4, 1.55, "cyan")]
    for t, d, x, w, col in side:
        box(s, x, 2.3, w, 1.7, [{"text": t, "size": 14, "bold": True}, {"text": d, "size": 12}], fill=col, edge=col, margin=0.08)
    ym = 3.0
    for x1, x2, lab in [(X0 + 1.55, X0 + 2.4, "MQTT"), (X0 + 3.95, X0 + 4.8, "subscribe"), (X0 + 8.55, X0 + 9.4, "WebSocket")]:
        line(s, x1 + 0.02, ym, x2 - 0.02, ym, color=INK, w=2)
        label(s, x1, ym - 0.32, x2 - x1, lab, size=12)
    for x1, x2, lab in [(X0 + 2.4, X0 + 1.55, "faults"), (X0 + 4.8, X0 + 3.95, "control"), (X0 + 9.4, X0 + 8.55, "REST")]:
        line(s, x1 - 0.02, 3.6, x2 + 0.02, 3.6, color=SLATE, w=1.5, dash=True)
        label(s, min(x1, x2), 3.64, abs(x2 - x1), lab, size=12)
    tb(s, X0 + 9.4, 4.15, 1.55, 0.5, "Chrome · Edge · Firefox", size=12, color=MUTED, align="c")
    for x, nm, d in [(bx, "InfluxDB v2", "telemetry_raw 30 d\ntelemetry_agg 1 y"), (bx + 1.95, "PostgreSQL 14", "machines · users · alerts\nmaintenance · feedback")]:
        line(s, x + 0.9, 5.2, x + 0.9, 5.8, color=INK, w=2)
        box(s, x, 5.8, 1.8, 1.05, [{"text": nm, "size": 14, "bold": True}] + [{"text": t, "size": 12} for t in d.split("\n")], fill="cyan", edge="cyan", margin=0.05)
    box(s, X0, 5.55, 4.4, 1.3, [{"text": "Two switches in .env", "size": 14, "bold": True},
                                {"text": "INGEST_MODE = direct | mqtt", "size": 13}, {"text": "STORAGE_MODE = memory | db", "size": 13},
                                {"text": "→ runs with no Docker and no database", "size": 13, "italic": True, "color": MUTED}], fill="grey", edge="white", align="l", margin=0.2)

    # ================================================================== 11  data flow: one 500 ms cycle
    s = new("Data Flow – One 500 ms Cycle",
            "Follow one reading. It is generated, published, validated, then OEE is updated, the ML model scores it, rules and trend are checked, alerts are raised or cleared, "
            "the data is stored and a snapshot is pushed to the browser. The whole loop is measured at about half a second at the 95th percentile - well under the 1 second target.")
    steps = [("Generate", "Simulator creates a reading per machine: temperature, vibration, status, cycles", "blue"),
             ("Publish", "Sent as JSON over MQTT / WebSockets through Mosquitto", "orange"),
             ("Validate", "Bad values are rejected at the gateway (FR-09)", "red"),
             ("Update OEE", "Rolling 1 h and 8 h windows (FR-02)", "green"),
             ("Score with ML", "Isolation Forest, 0 – 100 % (FR-04)", "yellow"),
             ("Rules + trend", "Hard limits → Red · What-If → Orange (FR-03, FR-07)", "orange"),
             ("Raise / clear alerts", "Priority Red > Orange > Yellow, debounced over 1 – 2 readings", "violet"),
             ("Store + push", "Batched write to the databases, snapshot to the dashboard", "cyan")]
    cs = cols(4, gap=0.55)
    ny, nh = [1.25, 3.55], 1.65
    for i, (t, d, col) in enumerate(steps):
        x, w = cs[i % 4]
        y = ny[i // 4]
        box(s, x, y, w, nh, "", fill=col, edge=col)
        badge(s, x + 0.35, y + 0.35, 0.42, str(i + 1), size=14)
        tb(s, x + 0.7, y + 0.1, w - 0.8, 0.5, t, size=15, bold=True, anchor="m")
        tb(s, x + 0.15, y + 0.72, w - 0.3, 0.9, d, size=13)
        if i % 4 < 3:
            line(s, x + w + 0.04, y + nh / 2, cs[i % 4 + 1][0] - 0.04, y + nh / 2, color=INK, w=2.25)
    xa = cs[3][0] + cs[3][1] / 2
    xb = cs[0][0] + cs[0][1] / 2
    ymid = (ny[0] + nh + ny[1]) / 2
    path(s, [(xa, ny[0] + nh + 0.02), (xa, ymid), (xb, ymid), (xb, ny[1] - 0.04)], color=INK, w=2.25)
    stats = [("500 ms", "the loop repeats every half second per machine"), ("0.51 s", "end-to-end update latency, 95th percentile (limit 1 s)"),
             ("≤ 25 ms", "worst-case ML inference measured (limit 200 ms)")]
    for (x, w), (n, d) in zip(cols(3), stats):
        box(s, x, 5.5, w, 1.3, "", fill="grey", edge="white")
        tb(s, x + 0.2, 5.58, w - 0.4, 0.6, n, size=28, bold=True, color=GREEN_D, anchor="m")
        tb(s, x + 0.2, 6.2, w - 0.4, 0.55, d, size=13)

    # ================================================================== 12  DFD level 0
    s = new("Data Flow Diagram – Level 0 (Context)",
            "The whole system is one process. External entities are the four user classes from the SRS and the virtual factory. Arrows show what data crosses the system boundary.")
    ocx, ocy, rx, ry = X0 + 5.475, 4.0, 1.5, 1.8
    ov = s.shapes.add_shape(MSO_SHAPE.OVAL, E(ocx - rx), E(ocy - ry), E(2 * rx), E(2 * ry))
    ov.shadow.inherit = False
    ov.fill.solid()
    ov.fill.fore_color.rgb = SOFT["green"]
    ov.line.color.rgb = EDGE["green"]
    ov.line.width = Pt(2.25)
    tf = ov.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    from deckkit import _fill_paras
    _fill_paras(tf, [{"text": "0", "size": 14, "bold": True, "color": MUTED}, {"text": "Real-Time OEE & Anomaly Dashboard", "size": 17, "bold": True}], 17, False, INK, "c", False)
    ew, eh = 2.15, 0.85
    left = [("Virtual Factory (machines)", 3.3, "blue"), ("System Administrator", 4.7, "violet")]
    right = [("Floor Manager", 2.6, "orange"), ("Maintenance Engineer", 4.0, "red"), ("Production Analyst", 5.4, "cyan")]
    for nm, yc, col in left:
        box(s, X0, yc - eh / 2, ew, eh, nm, fill=col, edge=col, size=15, bold=True, shape=MSO_SHAPE.RECTANGLE)
    for nm, yc, col in right:
        box(s, X1 - ew, yc - eh / 2, ew, eh, nm, fill=col, edge=col, size=15, bold=True, shape=MSO_SHAPE.RECTANGLE)

    def hflow(y, side_, text, to_system):
        xe = ocx - rx * math.sqrt(max(0.0, 1 - ((y - ocy) / ry) ** 2)) if side_ == "L" else ocx + rx * math.sqrt(max(0.0, 1 - ((y - ocy) / ry) ** 2))
        xs = X0 + ew if side_ == "L" else X1 - ew
        a, b = ((xs + 0.03, y), (xe - 0.03, y)) if side_ == "L" else ((xs - 0.03, y), (xe + 0.03, y))
        if (side_ == "L") != to_system:
            a, b = b, a
        line(s, a[0], a[1], b[0], b[1], color=INK, w=1.9)
        lo, hi = min(a[0], b[0]), max(a[0], b[0])
        tb(s, lo, y - 0.3, hi - lo, 0.26, text, size=12, color=MUTED, align="c", anchor="m")

    hflow(3.3, "L", "telemetry readings", True)
    hflow(4.7, "L", "configuration", True)
    hflow(2.35, "R", "live status · OEE · alerts", False)
    hflow(2.85, "R", "What-If request", True)
    hflow(3.75, "R", "alerts · live status", False)
    hflow(4.25, "R", "acknowledge · feedback", True)
    hflow(5.15, "R", "historical reports", False)
    hflow(5.65, "R", "report request", True)

    # ================================================================== 13  DFD level 1
    s = new("Data Flow Diagram – Level 1 (Seven Processes)",
            "The context process decomposed into seven processes P1 to P7 and three data stores. Telemetry enters through the MQTT broker, is validated, turned into OEE, "
            "checked by the detectors, and finally shown, stored and reported. Processes and stores follow the SRS section 9.4.")
    pw = (W - 6 * 0.22) / 7
    px = [X0 + i * (pw + 0.22) for i in range(7)]
    procs = [("P1", "Simulate telemetry", "blue"), ("P2", "Validate & ingest", "red"), ("P3", "Calculate OEE", "green"), ("P4", "Detect faults & anomalies", "yellow"),
             ("P5", "What-If projection", "orange"), ("P6", "Manage alerts & maintenance", "violet"), ("P7", "Display & report", "cyan")]
    py, ph = 2.85, 1.55
    for x, (pid, nm, col) in zip(px, procs):
        box(s, x, py, pw, ph, [{"text": pid, "size": 15, "bold": True, "color": RED}, {"text": nm, "size": 14, "bold": True}], fill=col, edge=col, margin=0.05)
    for i in range(6):
        line(s, px[i] + pw + 0.01, py + ph / 2, px[i + 1] - 0.01, py + ph / 2, color=INK, w=2)
    ey, eh2 = 1.1, 0.85
    ents = [(0, "System Administrator", "configure"), (4, "Floor Manager", "What-If request"), (5, "Maintenance Engineer", "acknowledge, notes"), (6, "Production Analyst", "report request")]
    for idx, nm, lab in ents:
        box(s, px[idx], ey, pw, eh2, nm, fill="white", edge="grey", size=13, bold=True, shape=MSO_SHAPE.RECTANGLE, margin=0.04)
        cx = px[idx] + 0.28
        line(s, cx, ey + eh2 + 0.02, cx, py - 0.03, color=INK, w=1.75, back_arrow=True)
        tb(s, cx + 0.1, ey + eh2 + 0.15, pw - 0.4, 0.7, lab, size=11, color=MUTED)
    sy, sh_ = 5.45, 0.85
    stores = [("D3  MQTT broker", "transient", px[0], px[1] + pw - px[0], "orange"), ("D1  InfluxDB", "time series", px[2], px[4] + pw - px[2], "cyan"),
              ("D2  PostgreSQL", "metadata · alerts · maintenance · feedback", px[5], px[6] + pw - px[5], "blue")]
    for nm, d, x, w, col in stores:
        box(s, x, sy, w, sh_, [{"text": nm, "size": 14, "bold": True}, {"text": d, "size": 11}], fill=col, edge=col, margin=0.04, shape=MSO_SHAPE.RECTANGLE)
    arrows = [(px[0] + 0.28, False, "publish", px[0]), (px[1] + 0.28, True, "subscribe", px[1]), (px[2] + 0.28, False, "OEE windows", px[2]), (px[5] + 0.28, None, "save / load", px[5]),
              (px[6] + 0.28, True, "history", px[6])]
    for ax_, up, lab, lx in arrows:
        if up is None:
            line(s, ax_, py + ph + 0.02, ax_, sy - 0.02, color=INK, w=1.75, back_arrow=True)
        elif up:
            line(s, ax_, sy - 0.02, ax_, py + ph + 0.02, color=INK, w=1.75)
        else:
            line(s, ax_, py + ph + 0.02, ax_, sy - 0.02, color=INK, w=1.75)
        tb(s, ax_ + 0.1, py + ph + 0.35, pw - 0.4, 0.5, lab, size=11, color=MUTED)
    tb(s, X0, 6.5, W, 0.4, "Arrows between processes: telemetry → valid readings → OEE → alerts → What-If / maintenance → dashboard and reports.  "
       "Process names are my decomposition of the SRS functions (§3.2).", size=12, color=MUTED, italic=True)

    # ================================================================== 14  methodology
    s = new("Methodology",
            "An incremental, requirements-driven approach: SRS first, then design, then build in ten milestones, verify, and deploy and document. "
            "Each milestone ended with working code and passing tests - that is the Definition of Done from SRS section 13.2.")
    phases = [("1  Requirements", "blue", ["SRS: 12 functional and 18 non-functional requirements", "6 user stories, 8 use cases", "Constraints C-01 to C-07"]),
              ("2  Design", "cyan", ["Architecture and data-flow diagrams", "UML: use case, class, sequence, activity", "Database schema"]),
              ("3  Build", "green", ["Ten incremental milestones (M1 – M10)", "Simulator → backend → ML → What-If → dashboard", "Docker Compose, scripts"]),
              ("4  Verify", "orange", ["45 automated tests", "13-check end-to-end smoke test", "Accessibility, 3 browsers, live databases"]),
              ("5  Document", "violet", ["Project report and this presentation", "README, setup guides", "Auto-generated API docs at /docs"])]
    cs5 = cols(5, gap=0.2)
    for (x, w), (t, col, bl) in zip(cs5, phases):
        shp = s.shapes.add_shape(MSO_SHAPE.CHEVRON, E(x), E(1.2), E(w + 0.22), E(0.95))
        shp.shadow.inherit = False
        shp.fill.solid()
        shp.fill.fore_color.rgb = EDGE[col]
        shp.line.fill.background()
        shp.adjustments[0] = 0.28
        tf = shp.text_frame
        tf.margin_left, tf.margin_right = E(0.28), E(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _fill_paras = __import__("deckkit")._fill_paras
        _fill_paras(tf, [{"text": t, "size": 13, "bold": True}], 13, True, WHITE, "l", False)
        box(s, x, 2.4, w, 3.35, "", fill=col, edge=col)
        tb(s, x + 0.15, 2.55, w - 0.3, 3.1, [{"text": b, "size": 15, "bullet": True, "space_after": 10} for b in bl], size=15)
    box(s, X0, 6.0, W, 0.8, "Every milestone ended with working code and passing tests – the Definition of Done in SRS §13.2", fill="grey", edge="white", size=16, bold=True, color=GREEN_D)

    # ================================================================== 15  modules
    s = new("Modules (SRS §3.2)",
            "The eight modules from the SRS, each with the file where it lives, so every module can be shown in the code if asked.")
    mods8 = [("Virtual Factory Simulator", "Emulates 5–20 machines; emits temperature, vibration, status and cycles every 500 ms.", "simulator/core.py", "blue"),
             ("OEE Calculation Engine", "Availability × Performance × Quality over rolling 1 h and 8 h windows.", "backend/app/services/oee.py", "green"),
             ("Rule-Based Alert Engine", "Deterministic hard limits raise Red alerts instantly.", "backend/app/services/rules.py", "red"),
             ("ML Anomaly Scorer", "Pre-trained Isolation Forest returns a 0–100 % anomaly score.", "backend/app/ml/scorer.py", "yellow"),
             ("Interactive 2D Dashboard", "Machine grid, live charts, alert feed, ML gauges, reports, controls.", "frontend/src/pages/", "cyan"),
             ("Maintenance & Feedback Logger", "Acknowledge, assign technician, notes, resolve, ML feedback.", "backend/app/services/alerts.py", "violet"),
             ("What-If Degradation Simulator", "Injects a temperature rise and projects the crossing with y = mx + c.", "backend/app/services/whatif.py", "orange"),
             ("Historical Reporting API", "7-day OEE, downtime per shift and ML accuracy.", "backend/app/api/routes.py", "grey")]
    for i, (t, d, f, col) in enumerate(mods8):
        (x, w), y = cols(4)[i % 4], 1.2 + (i // 4) * 2.85
        box(s, x, y, w, 2.6, "", fill=col, edge=col)
        tb(s, x + 0.15, y + 0.12, w - 0.3, 0.7, t, size=16, bold=True)
        tb(s, x + 0.15, y + 0.85, w - 0.3, 1.25, d, size=13)
        tb(s, x + 0.15, y + 2.2, w - 0.3, 0.3, f, size=11, color=MUTED, italic=True)

    # ================================================================== 16  technology stack
    s = new("Technologies Used",
            "Each layer and the tool chosen for it, all taken from the SRS technology stack. Everything is open source.")
    rows = [("Data source", ["Python simulator", "paho-mqtt", "NumPy"], "blue"), ("Messaging", ["Eclipse Mosquitto 2", "MQTT over WebSockets"], "orange"),
            ("Backend", ["FastAPI", "Pydantic v2", "asyncio", "WebSocket", "OpenAPI / Swagger"], "green"),
            ("Intelligence", ["scikit-learn Isolation Forest", "NumPy", "Least-squares regression"], "yellow"),
            ("Storage", ["InfluxDB v2 (time series)", "PostgreSQL 14", "SQLAlchemy"], "cyan"),
            ("Frontend", ["React 18", "TypeScript", "Vite", "Chart.js"], "violet"),
            ("Deploy & quality", ["Docker Compose", "nginx", "pytest", "Vitest", "Playwright", "axe-core"], "grey")]
    rh, rg = 0.7, (5.8 - 7 * 0.7) / 6
    for i, (layer, chips, col) in enumerate(rows):
        y = 1.1 + i * (rh + rg)
        box(s, X0, y, 2.3, rh, layer, fill=col, edge=col, size=16, bold=True)
        x = X0 + 2.5
        for c in chips:
            cw = len(c) * 0.092 + 0.45
            box(s, x, y + 0.08, cw, rh - 0.16, c, fill="white", edge=col, size=14, line=1.75, margin=0.05)
            x += cw + 0.15

    # ================================================================== 17  OEE
    s = new("OEE Calculation – Live Values",
            "Four rings: plant OEE and its three factors, measured in the end-to-end test over the rolling one-hour window on five machines. "
            "Reading the rings tells you WHY OEE is low: here performance is the weakest factor. Bottom left explains how it is computed efficiently; bottom right the measured figures.")
    rings = [("OEE", 83.1, "overall effectiveness", RGBColor(0x22, 0xC5, 0x5E)), ("Availability", 97.9, "run time ÷ planned time", RGBColor(0x25, 0x63, 0xEB)),
             ("Performance", 86.1, "actual ÷ ideal speed", RGBColor(0x08, 0x91, 0xB2)), ("Quality", 98.6, "good ÷ all parts", RGBColor(0x7C, 0x3A, 0xED))]
    for (x, w), (nm, v, sub, colr) in zip(cols(4), rings):
        d = 2.2
        cd = CategoryChartData()
        cd.categories = ["value", "rest"]
        cd.add_series(nm, (v, 100 - v))
        gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, E(x + (w - d) / 2), E(1.1), E(d), E(d), cd)
        ch = gf.chart
        ch.has_legend = False
        ch.has_title = False
        plot = ch.plots[0]
        plot.has_data_labels = False
        hs = plot._element.find(qn("c:holeSize"))
        if hs is None:
            hs = etree.SubElement(plot._element, qn("c:holeSize"))
        hs.set("val", "72")
        ser = plot.series[0]
        for idx, colr_ in enumerate((colr, RGBColor(0xE2, 0xE8, 0xF0))):
            pt = ser.points[idx]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = colr_
            pt.format.line.fill.background()
        tb(s, x + (w - d) / 2, 1.1 + d / 2 - 0.35, d, 0.7, f"{v:.1f} %", size=26, bold=True, color=INK, align="c", anchor="m")
        tb(s, x, 3.4, w, 0.35, nm, size=17, bold=True, align="c")
        tb(s, x, 3.77, w, 0.3, sub, size=13, color=MUTED, align="c")
    card(s, X0, 4.4, 5.4, 2.4, "blue", "How it is computed", ["Rolling 1 h and 8 h windows built from 5-second buckets – constant-time reads, so 20 machines are no problem",
                                                                "Plant OEE uses the summed run time, cycles and good parts of all machines",
                                                                "Performance is capped at 100 %; cycle-counter resets are handled"], size=14)
    card(s, X0 + 5.6, 4.4, W - 5.6, 2.4, "green", "Measured in the end-to-end test", ["1 h window, 5 machines:  A 97.9 % · P 86.1 % · Q 98.6 %  →  OEE 83.1 %",
                                                                                      "8 h window: OEE 83.3 %",
                                                                                      "Formula verified exactly by a unit test (test_oee_formula_exact)"], size=14)

    # ================================================================== 18  alert decision flowchart
    s = new("Alert Decision Flow – Every 500 ms per Machine",
            "This is the heart of the system. Each reading goes through three questions in strict priority. Red is certain and immediate. Orange uses a regression line to predict a crossing within "
            "60 seconds. Yellow comes from the Isolation Forest when the score stays above 85 percent. Higher levels suppress lower ones, and every alert is handed to the alert manager.")
    dx, dw = X0, 4.7
    cx = dx + dw / 2
    box(s, cx - 1.8, 1.1, 3.6, 0.5, "Telemetry reading received and validated", fill="blue", edge="blue", size=14)
    qs = [("Temperature or vibration above its limit, or status STOPPED?", 1.8, "red", "RED", "Hard fault – rule engine, immediate", "FR-03"),
          ("Trend line reaches the limit within 60 s? (R² ≥ 0.6)", 3.2, "orange", "ORANGE", "Projected fault – with ETA countdown", "FR-07"),
          ("ML anomaly score above 85 % for 2 readings?", 4.6, "yellow", "YELLOW", "ML caution – Isolation Forest", "FR-04")]
    line(s, cx, 1.62, cx, 1.78, color=INK, w=2)
    for q, y, col, nm, d, fr in qs:
        dm = s.shapes.add_shape(MSO_SHAPE.FLOWCHART_DECISION, E(dx), E(y), E(dw), E(1.15))
        dm.shadow.inherit = False
        dm.fill.solid()
        dm.fill.fore_color.rgb = SOFT["grey"]
        dm.line.color.rgb = EDGE["grey"]
        dm.line.width = Pt(2)
        tb(s, dx + 0.85, y + 0.17, dw - 1.7, 0.81, q, size=13, align="c", anchor="m")
        line(s, dx + dw + 0.02, y + 0.575, dx + dw + 0.78, y + 0.575, color=INK, w=2)
        label(s, dx + dw + 0.02, y + 0.25, 0.8, "yes", size=12, bold=True)
        box(s, dx + dw + 0.8, y + 0.125, 3.05, 0.9, [{"text": nm, "size": 16, "bold": True}, {"text": d, "size": 12}], fill=col, edge=col, line=2, margin=0.08)
        line(s, dx + dw + 3.87, y + 0.575, dx + dw + 4.37, y + 0.575, color=SLATE, w=1.5, dash=True)
    for y in (2.95, 4.35):
        line(s, cx, y + 0.01, cx, y + 0.24, color=INK, w=2)
        label(s, cx + 0.08, y - 0.01, 0.5, "no", size=12, bold=True, align="l")
    line(s, cx, 5.76, cx, 5.98, color=INK, w=2)
    label(s, cx + 0.08, 5.74, 0.5, "no", size=12, bold=True, align="l")
    box(s, cx - 1.8, 6.0, 3.6, 0.5, "GREEN – normal operation", fill="green", edge="green", size=14, bold=True)
    amx = dx + dw + 4.4
    box(s, amx, 1.8 + 0.125, X1 - amx, (4.6 + 1.15 - 1.8 - 0.25), [{"text": "Alert manager", "size": 15, "bold": True}, {"text": "acknowledge", "size": 13}, {"text": "assign", "size": 13},
                                                                    {"text": "resolve", "size": 13}, {"text": "ML feedback", "size": 13}], fill="violet", edge="violet", margin=0.05)
    tb(s, dx, 6.6, W, 0.3, "Higher levels suppress lower ones  ·  an alert clears when its condition is absent for 3 – 4 s", size=13, color=MUTED, italic=True)

    # ================================================================== 19  four alert states
    s = new("Four Machine States",
            "A reference table for the four colours. Each has a trigger, a detector and what the operator should do. Colour is never the only cue: every state also has an icon and a text label "
            "(a usability and accessibility requirement).")
    cw = [2.2, 3.7, 2.15, 2.6]
    cxs = [X0]
    for w_ in cw[:-1]:
        cxs.append(cxs[-1] + w_ + 0.1)
    heads = ["State", "Trigger", "Detected by", "Operator sees / does"]
    for x, w_, h in zip(cxs, cw, heads):
        box(s, x, 1.15, w_, 0.5, h, fill="red", solid=RED, edge="red", size=15, bold=True, color=WHITE, align="l", margin=0.15)
    rows4 = [("green", "✓  GREEN\nNormal", "No limit crossed, no anomaly, no projected crossing", "–", "Tile stays green"),
             ("yellow", "▲  YELLOW\nML caution", "Anomaly score above 85 % for 2 readings – still below every hard limit", "Isolation Forest\nFR-04", "Early warning: inspect, then give true / false-positive feedback"),
             ("orange", "↗  ORANGE\nProjected fault", "Trend line y = m·x + c reaches the limit within 60 s", "What-If regression\nFR-07", "ETA countdown: act before the breach"),
             ("red", "✕  RED\nHard fault", "Temperature or vibration above its limit, or machine STOPPED", "Rule engine\nFR-03", "Immediate alert: acknowledge, assign, resolve")]
    for i, (col, st, trig, by, act) in enumerate(rows4):
        y = 1.8 + i * 1.22
        head, sub = st.split("\n")
        box(s, cxs[0], y, cw[0], 1.1, [{"text": head, "size": 17, "bold": True}, {"text": sub, "size": 14}], fill=col, edge=col, line=2, align="l", margin=0.15)
        for x, w_, t in zip(cxs[1:], cw[1:], (trig, by, act)):
            box(s, x, y, w_, 1.1, t.split("\n"), fill="grey", edge="white", size=14, align="l", margin=0.15)
    tb(s, X0, 6.7, W, 0.3, "Priority  Red > Orange > Yellow > Green  ·  also shown with an icon and a label, never by colour alone (NFR-U1, NFR-U2)", size=13, color=MUTED, italic=True)

    # ================================================================== 20  ML pipeline
    s = new("ML Anomaly Scoring – Isolation Forest",
            "Top lane: done once, offline. A synthetic dataset of normal behaviour trains an Isolation Forest; the raw model output is calibrated to a 0 to 100 percent score and saved as model.pkl. "
            "Bottom lane: every 500 ms the live reading is converted to the same features and scored; two readings above 85 percent raise a Yellow caution. "
            "Validated on injected faults: six of six detected, no false alarms. If the ML service fails, the last score is held and the red rules keep working.")
    lanes = [("Offline\n(once)", 1.2, [("Synthetic normal data", "68 687 rows · 20 machine profiles", "blue"), ("Feature engineering", "6 normalised features", "cyan"),
                                       ("Train Isolation Forest", "unsupervised – no fault labels needed", "violet"), ("Calibrate score", "mapped to 0–100 %, alert threshold 85 %", "yellow"),
                                       ("Save model.pkl", "loaded at start-up (C-04: trained once)", "grey")]),
             ("Online\n(every 500 ms)", 3.15, [("Live reading", "temperature, vibration, status, cycles", "blue"), ("Same features", "relative to the machine's own profile", "cyan"),
                                            ("Score the reading", "≤ 25 ms measured (limit 200 ms)", "violet"), ("Above 85 % for 2 readings?", "and below every hard limit", "yellow"),
                                            ("YELLOW caution", "operator gives true / false feedback", "orange")])]
    bw, bg = 1.55, 0.35
    for lab, y, items in lanes:
        box(s, X0, y, 1.3, 1.6, [{"text": t, "size": 14, "bold": True} for t in lab.split("\n")], fill="grey", edge="white")
        for i, (t, d, col) in enumerate(items):
            x = X0 + 1.5 + i * (bw + bg)
            box(s, x, y, bw, 1.6, [{"text": t, "size": 13, "bold": True, "space_after": 3}, {"text": d, "size": 12}], fill=col, edge=col, margin=0.07)
            if i < 4:
                line(s, x + bw + 0.02, y + 0.8, x + bw + bg - 0.02, y + 0.8, color=INK, w=2)
    for (x, w), (n, d) in zip(cols(4), [("6 / 6", "injected anomalies detected"), ("100 %", "precision on injected faults"),
                                       ("0", "false alarms per hour on healthy machines"), ("≤ 25 ms", "worst-case inference (limit 200 ms)")]):
        box(s, x, 5.0, w, 1.45, "", fill="green", edge="green")
        tb(s, x + 0.15, 5.08, w - 0.3, 0.65, n, size=30, bold=True, color=GREEN_D, anchor="m")
        tb(s, x + 0.15, 5.78, w - 0.3, 0.6, d, size=13)
    tb(s, X0, 6.55, W, 0.3, "If the ML service fails, the last score is kept and the Red rules keep working (NFR-R1).", size=13, color=MUTED, italic=True)

    # ================================================================== 21  What-If
    s = new("What-If Degradation & Projection (FR-07)",
            "The manager injects a controlled temperature rise over three minutes. The backend fits a straight line y = m x + c through the last three minutes of readings and predicts "
            "when the hard limit will be crossed. If that is within 60 seconds and the fit is good, the machine turns Orange with a countdown. The chart is an illustration with a slope of 0.1 degrees per second.")
    pts = [(t, 62 + 0.1 * (t + 180) + 0.7 * math.sin(t / 6.0)) for t in range(-180, 1, 10)]
    cd = XyChartData()
    s1 = cd.add_series("Temperature readings")
    for t, v in pts:
        s1.add_data_point(t, round(v, 2))
    s2 = cd.add_series("Fitted line  y = m·x + c")
    s2.add_data_point(-180, 62.0), s2.add_data_point(60, 86.0)
    s3 = cd.add_series("Hard limit 85 °C")
    s3.add_data_point(-180, 85), s3.add_data_point(60, 85)
    s4 = cd.add_series("Now")
    s4.add_data_point(0, 55), s4.add_data_point(0, 95)
    s5 = cd.add_series("Predicted crossing (ETA ≈ 50 s)")
    s5.add_data_point(50, 85)
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, E(X0), E(1.1), E(6.6), E(4.95), cd)
    ch = gf.chart
    ch.font.size, ch.font.name = Pt(12), FONT
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(11)
    sers = ch.plots[0].series
    sers[0].format.line.fill.background()
    sers[0].marker.style, sers[0].marker.size = XL_MARKER_STYLE.CIRCLE, 7
    sers[0].marker.format.fill.solid()
    sers[0].marker.format.fill.fore_color.rgb = EDGE["blue"]
    sers[0].marker.format.line.fill.background()
    for ser, colr, wd, dash in ((sers[1], EDGE["orange"], 2.5, False), (sers[2], EDGE["red"], 2.5, False), (sers[3], SLATE, 1.5, True)):
        ser.marker.style = XL_MARKER_STYLE.NONE
        ser.format.line.color.rgb = colr
        ser.format.line.width = Pt(wd)
        ser.smooth = False
        if dash:
            ser.format.line.dash_style = MSO_LINE.DASH
    sers[4].format.line.fill.background()
    sers[4].marker.style, sers[4].marker.size = XL_MARKER_STYLE.DIAMOND, 12
    sers[4].marker.format.fill.solid()
    sers[4].marker.format.fill.fore_color.rgb = EDGE["red"]
    xa, ya = ch.category_axis, ch.value_axis
    xa.crosses = XL_AXIS_CROSSES.MINIMUM
    ya.crosses = XL_AXIS_CROSSES.MINIMUM
    xa.minimum_scale, xa.maximum_scale, xa.major_unit = -180, 60, 30
    ya.minimum_scale, ya.maximum_scale, ya.major_unit = 55, 95, 5
    for ax, ttl in ((xa, "Time relative to now (s)"), (ya, "Temperature (°C)")):
        ax.has_title = True
        ax.axis_title.text_frame.text = ttl
        ax.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        ax.axis_title.text_frame.paragraphs[0].runs[0].font.bold = False
        ax.tick_labels.font.size = Pt(12)
        if ax.has_major_gridlines:
            ax.major_gridlines.format.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    tb(s, X0, 6.15, 6.6, 0.55, "Illustrative example: the line predicts the 85 °C limit in about 50 s (≤ 60 s), so the machine turns Orange.", size=13, color=MUTED, italic=True)
    rx0, rw0 = X0 + 6.85, W - 6.85
    st = ["Fit a straight line through the last **3 minutes** of temperature", "Compute **ETA = (limit − value now) ÷ slope**", "If ETA ≤ **60 s** and the fit is good (R² ≥ 0.6, slope ≥ 0.03 °C/s) → **ORANGE** + countdown",
          "If the limit is really crossed → **RED**"]
    for i, t in enumerate(st):
        y = 1.1 + i * 1.12
        box(s, rx0, y, rw0, 0.98, "", fill="orange", edge="orange")
        badge(s, rx0 + 0.42, y + 0.49, 0.46, str(i + 1), fill=EDGE["orange"])
        tb(s, rx0 + 0.85, y, rw0 - 1.0, 0.98, t, size=14, anchor="m")
    box(s, rx0, 5.7, rw0, 1.15, [{"text": "Measured in tests", "size": 14, "bold": True}, {"text": "Orange raised with ETA ≈ 28 s, then Red at 85.3 °C. Flat noise and short windows do not trigger it (unit-tested).", "size": 13}],
        fill="green", edge="green", align="l", margin=0.18)

    # ================================================================== 22  use case diagram
    s = new("UML – Use Case Diagram (SRS §7.2)",
            "Eight use cases from the SRS catalogue and four actors. The Maintenance Engineer reacts to alerts; the Floor Manager watches and plans with What-If; the Production Analyst reads reports; "
            "the System Administrator deploys and scales. A login is required for the write operations (acknowledge, log maintenance, What-If).")
    bxl, bxw = X0 + 1.85, 7.25
    box(s, bxl, 1.1, bxw, 5.75, "", fill=None, edge="grey", radius=0.08, line=2)
    tb(s, bxl + 0.2, 1.17, bxw - 0.4, 0.35, "Real-Time OEE and Anomaly Dashboard for a Smart Factory System", size=13, bold=True, color=MUTED)
    ucs = {"ack": ("UC-01  Acknowledge an Alarm", bxl + 0.3, 1.85), "alerts": ("UC-02  Receive Real-Time Alerts", bxl + 0.3, 2.85), "log": ("UC-03  Log Maintenance Activity", bxl + 0.3, 3.85),
           "dash": ("UC-06  View Real-Time Dashboard", bxl + 0.3, 4.85), "whatif": ("UC-07  Trigger “What-If” Scenario", bxl + 0.3, 5.85),
           "rep": ("UC-04  View Historical Reports", bxl + 3.75, 2.35), "dep": ("UC-05  Deploy System", bxl + 3.75, 4.35), "scale": ("UC-08  Scale Simulated Machines", bxl + 3.75, 5.35)}
    ow, oh = 3.2, 0.7
    for k, (t, x, y) in ucs.items():
        o = box(s, x, y, ow, oh, t, fill="violet", edge="violet", size=13, shape=MSO_SHAPE.OVAL, margin=0.12)
    def left_of(k):
        t, x, y = ucs[k]
        return x, y + oh / 2
    def right_of(k):
        t, x, y = ucs[k]
        return x + ow, y + oh / 2
    actor(s, X0 + 0.7, 2.0, "Maintenance Engineer", label_w=1.6)
    actor(s, X0 + 0.7, 4.7, "Floor Manager", label_w=1.6)
    for k in ("ack", "alerts", "log", "dash"):
        line(s, X0 + 1.05, 2.55, *left_of(k), color=SLATE, w=1.5, arrow=False)
    for k in ("dash", "whatif"):
        line(s, X0 + 1.05, 5.25, *left_of(k), color=SLATE, w=1.5, arrow=False)
    actor(s, X1 - 0.7, 1.9, "Production Analyst", label_w=1.6)
    actor(s, X1 - 0.7, 4.7, "System Administrator", label_w=1.6)
    line(s, X1 - 1.05, 2.45, *right_of("rep"), color=SLATE, w=1.5, arrow=False)
    for k in ("dep", "scale"):
        line(s, X1 - 1.05, 5.25, *right_of(k), color=SLATE, w=1.5, arrow=False)
    tb(s, bxl + 3.75, 1.75, ow, 0.45, "Write operations (UC-01, 03, 07) require login", size=12, color=MUTED, italic=True, align="c")

    # ================================================================== 23  class diagram
    s = new("UML – Class Diagram: 12 Domain Classes (SRS §8.3)",
            "The twelve domain classes from the SRS, with the main attributes and operations. Row one is the data path (Simulator, Machine, TelemetryData, OEEEngine). Row two is detection: "
            "the ML scorer and the rule engine raise Alerts; OEEEngine produces OEE. Row three is the people side: maintenance records belong to alerts, feedback rates alerts.")
    cs4 = cols(4, gap=0.3)
    ch_, pitch = 1.6, 1.9
    cls = {
        "Simulator": (0, 0, "blue", ["machineCount: int", "emitIntervalMs: int", "isRunning: bool"], ["generateTelemetry()", "injectDegradation(id)"]),
        "Machine": (1, 0, "blue", ["machineID: UUID", "status: enum", "oeeScore: float"], ["updateTelemetry()", "setStatus(status)"]),
        "TelemetryData": (2, 0, "green", ["machineId · timestamp", "temperature: float", "vibration: float"], ["validate(): bool", "toInfluxPoint()"]),
        "OEEEngine": (3, 0, "green", ["windowSize: int", "plannedProductionTime", "idealCycleTime: float"], ["calculateOEE(id, window)", "getTrend(id, days)"]),
        "MLAnomalyScorer": (0, 1, "yellow", ["modelPath: string", "threshold: float", "isLoaded: bool"], ["loadModel()", "getAnomalyScore(): float"]),
        "Alert": (1, 1, "orange", ["alertType · severity: enum", "isAcknowledged: bool", "isResolved: bool"], ["acknowledge(userId)", "resolve()"]),
        "RuleEngine": (2, 1, "red", ["ruleId: UUID", "rules: dict", "isEnabled: bool"], ["evaluateTelemetry(data)", "triggerAlert(rule, id)"]),
        "OEE": (3, 1, "green", ["availability: float", "performance: float", "quality · oeeScore: float"], ["calculateOEE()", "toReport(): dict"]),
        "User": (0, 2, "grey", ["username: string", "role: enum", "passwordHash: string"], ["authenticate(password)", "hasPermission(action)"]),
        "MaintenanceRecord": (1, 2, "violet", ["technicianName: string", "rootCause · resolutionNotes", "timeTaken: int"], ["addNote(note)", "markComplete()"]),
        "MLAccuracyFeedback": (2, 2, "violet", ["isTruePositive: bool", "isFalsePositive: bool", "providedBy: UUID"], ["isAccurate(): bool", "toReport(): dict"]),
        "DashboardConfig": (3, 2, "grey", ["machineCount: int", "thresholdTemperature", "anomalyThreshold: float"], ["loadConfig()", "validate(): bool"]),
    }
    pos = {}
    for name, (c, r, col, at, op) in cls.items():
        x, w = cs4[c]
        y = 1.1 + r * pitch
        class_box(s, x, y, w, ch_, name, at, op, col)
        pos[name] = (x, y, w, ch_)

    def mid(n, side):
        x, y, w, h = pos[n]
        return {"l": (x, y + h / 2), "r": (x + w, y + h / 2), "t": (x + w / 2, y), "b": (x + w / 2, y + h)}[side]
    line(s, *mid("Simulator", "r"), *mid("Machine", "l"), color=INK, w=2)
    line(s, *mid("Machine", "r"), *mid("TelemetryData", "l"), color=INK, w=2)
    line(s, *mid("OEEEngine", "l"), *mid("TelemetryData", "r"), color=INK, w=2, dash=True)
    line(s, *mid("Machine", "b"), *mid("Alert", "t"), color=INK, w=2)
    line(s, *mid("MLAnomalyScorer", "r"), *mid("Alert", "l"), color=INK, w=2, dash=True)
    line(s, *mid("RuleEngine", "l"), *mid("Alert", "r"), color=INK, w=2, dash=True)
    line(s, *mid("OEEEngine", "b"), *mid("OEE", "t"), color=INK, w=2, dash=True)
    line(s, *mid("Alert", "b"), *mid("MaintenanceRecord", "t"), color=INK, w=2, arrow=False)
    dmx, dmy = mid("Alert", "b")
    dm = s.shapes.add_shape(MSO_SHAPE.DIAMOND, E(dmx - 0.1), E(dmy), E(0.2), E(0.2))
    dm.shadow.inherit = False
    dm.fill.solid()
    dm.fill.fore_color.rgb = INK
    line(s, *mid("User", "r"), *mid("MaintenanceRecord", "l"), color=INK, w=2, arrow=False)
    ax, ay = pos["Alert"][0] + pos["Alert"][2] - 0.05, pos["Alert"][1] + pos["Alert"][3]
    fx, fy = pos["MLAccuracyFeedback"][0] + 0.05, pos["MLAccuracyFeedback"][1]
    line(s, ax, ay, fx, fy, color=INK, w=2, arrow=False)
    tb(s, X0, 6.6, W, 0.25, "Solid = association · dashed = dependency · ◆ = composition (an Alert owns its maintenance records) · also: Machine 1—0..* OEE, User 1—0..* feedback",
       size=11, color=MUTED, italic=True)

    # ================================================================== 24  sequence diagram
    s = new("UML – Sequence Diagram (SRS §9.2)",
            "Login first, then a loop that repeats every 500 ms: the simulator publishes, the backend validates and calculates OEE, asks the ML service for a score, saves telemetry and alerts, and pushes a live "
            "update to the dashboard. At the bottom an authenticated user acknowledges an alert and logs maintenance; the alert is marked In Progress.")
    names = ["User", "Simulator", "Backend API", "ML Service", "Database", "Dashboard UI"]
    lw = 1.5
    lxs = [X0 + lw / 2 + i * (W - lw) / 5 for i in range(6)]
    for x, nm in zip(lxs, names):
        box(s, x - lw / 2, 1.1, lw, 0.5, nm, fill="grey", edge="grey", size=14, bold=True, shape=MSO_SHAPE.RECTANGLE, margin=0.03)
        line(s, x, 1.62, x, 6.85, color=SLATE, w=1, arrow=False, dash=True)
    U, Sm, B, M, D, UI = lxs
    pitch = 0.29
    y = 1.95
    items = [("msg", U, UI, "1  Open dashboard", False), ("msg", UI, B, "2  Login", False), ("msg", B, D, "3  Verify credentials", False), ("msg", B, UI, "4  Dashboard unlocked", True),
             ("band", "Steps 5 – 10 repeat every 500 ms"), ("msg", Sm, B, "5  Publish telemetry", False), ("note", "6  Validate payload · calculate OEE · apply rules"),
             ("msg", B, M, "7  Request score", False), ("msg", M, B, "8  Score 0–100 %", True), ("msg", B, D, "9  Save telemetry and alerts", False),
             ("msg", B, UI, "10  Push live update", False), ("band", "User write operations (authenticated)"), ("msg", U, UI, "11  Acknowledge alert", False),
             ("msg", UI, B, "12  Update alert status", False), ("msg", B, D, "13  Log maintenance", False), ("msg", B, UI, "14  Alert marked “In Progress”", True)]
    for it in items:
        if it[0] == "msg":
            _, a, b, t, ret = it
            yy = y + pitch - 0.05
            line(s, a, yy, b + (0.04 if b < a else -0.04), yy, color=EDGE["blue"] if not ret else SLATE, w=1.5, dash=ret)
            _white_label(s, (a + b) / 2, y + 0.11, min(abs(b - a) - 0.1, len(t) * 0.082 + 0.25), t, size=12, h=0.21)
        elif it[0] == "band":
            box(s, lxs[0] - 0.1, y + 0.02, W - lw + 0.2, 0.25, it[1], fill="grey", edge="grey", size=12, bold=True, shape=MSO_SHAPE.RECTANGLE, margin=0.03)
        else:
            box(s, B - 2.3, y + 0.01, 4.6, 0.27, it[1], fill="yellow", edge="yellow", size=12, shape=MSO_SHAPE.RECTANGLE, margin=0.03)
        y += pitch

    # ================================================================== 25  activity diagram
    s = new("UML – Activity Diagram (SRS §9.3)",
            "The combined activity flow from the SRS. Data is monitored and checked against the hard rules first: a hard fault is a Red alert. Otherwise the ML anomaly detector decides between a Yellow "
            "caution and normal status. Every outcome is displayed; the engineer views the alert and either starts maintenance or carries on monitoring.")
    def act(x, y, w, h, t, col="blue"):
        return box(s, x, y, w, h, t, fill=col, edge=col, size=13, margin=0.06)

    def dec(x, y, w, h, t):
        d = s.shapes.add_shape(MSO_SHAPE.FLOWCHART_DECISION, E(x), E(y), E(w), E(h))
        d.shadow.inherit = False
        d.fill.solid()
        d.fill.fore_color.rgb = SOFT["grey"]
        d.line.color.rgb, d.line.width = EDGE["grey"], Pt(2)
        tb(s, x + 0.2, y + h * 0.2, w - 0.4, h * 0.6, t, size=12, align="c", anchor="m")

    def dot(cx, cy, end=False):
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, E(cx - 0.15), E(cy - 0.15), E(0.3), E(0.3))
        o.shadow.inherit = False
        o.fill.solid()
        o.fill.fore_color.rgb = INK
        o.line.fill.background()
        if end:
            o.fill.fore_color.rgb = WHITE
            o.line.color.rgb, o.line.width = INK, Pt(3)
    r1 = 1.75
    dot(X0 + 0.15, r1)
    act(X0 + 0.5, r1 - 0.35, 1.75, 0.7, "Monitor dashboard")
    act(X0 + 2.5, r1 - 0.35, 1.95, 0.7, "Receive machine data")
    act(X0 + 4.7, r1 - 0.35, 2.0, 0.7, "Apply hard-fault rules")
    dec(X0 + 6.95, r1 - 0.5, 1.7, 1.0, "Hard fault?")
    act(X0 + 9.0, r1 - 0.35, 1.95, 0.7, "Generate RED alert", "red")
    for a, b in [(X0 + 0.3, X0 + 0.5), (X0 + 2.25, X0 + 2.5), (X0 + 4.45, X0 + 4.7), (X0 + 6.7, X0 + 6.95)]:
        line(s, a, r1, b, r1, color=INK, w=1.75)
    line(s, X0 + 8.65, r1, X0 + 9.0, r1, color=INK, w=1.75)
    label(s, X0 + 8.55, r1 - 0.3, 0.5, "yes", size=12, bold=True)
    r2 = 3.4
    cx1 = X0 + 7.8
    line(s, cx1, r1 + 0.5, cx1, r2 - 0.35, color=INK, w=1.75)
    label(s, cx1 + 0.08, r1 + 0.62, 0.5, "no", size=12, bold=True, align="l")
    act(X0 + 6.8, r2 - 0.35, 2.0, 0.7, "Run ML anomaly detection", "yellow")
    dec(X0 + 3.6, r2 - 0.5, 1.9, 1.0, "Score above 85 %?")
    act(X0 + 0.45, r2 - 0.35, 2.2, 0.7, "Generate YELLOW caution", "yellow")
    line(s, X0 + 6.8, r2, X0 + 5.5, r2, color=INK, w=1.75)
    line(s, X0 + 3.6, r2, X0 + 2.65, r2, color=INK, w=1.75)
    label(s, X0 + 2.65, r2 - 0.32, 0.95, "yes", size=12, bold=True)
    by = 4.55
    box(s, X0 + 0.45, by, 10.5, 0.55, "Display alert or status  (Green · Yellow · Red)", fill="green", edge="green", size=15, bold=True)
    line(s, X0 + 1.55, r2 + 0.35, X0 + 1.55, by - 0.01, color=INK, w=1.75)
    line(s, X0 + 4.55, r2 + 0.5, X0 + 4.55, by - 0.01, color=INK, w=1.75)
    label(s, X0 + 4.63, r2 + 0.62, 1.2, "no: normal", size=12, bold=True, align="l")
    line(s, X0 + 9.975, r1 + 0.35, X0 + 9.975, by - 0.01, color=INK, w=1.75)
    r4 = 5.85
    line(s, X0 + 1.6, by + 0.55, X0 + 1.6, r4 - 0.35, color=INK, w=1.75)
    act(X0 + 0.45, r4 - 0.35, 2.3, 0.7, "Engineer views alert")
    dec(X0 + 3.2, r4 - 0.5, 2.0, 1.0, "Action needed?")
    act(X0 + 5.9, r4 - 0.35, 2.2, 0.7, "Start maintenance", "violet")
    act(X0 + 8.6, r4 - 0.35, 1.8, 0.7, "Alert handled", "green")
    for a, b in [(X0 + 2.75, X0 + 3.2), (X0 + 5.2, X0 + 5.9), (X0 + 8.1, X0 + 8.6)]:
        line(s, a, r4, b, r4, color=INK, w=1.75)
    label(s, X0 + 5.15, r4 - 0.32, 0.8, "yes", size=12, bold=True)
    line(s, X0 + 10.4, r4, X0 + 10.65, r4, color=INK, w=1.75)
    dot(X0 + 10.8, r4, end=True)
    path(s, [(X0 + 4.2, r4 + 0.5), (X0 + 4.2, 6.68), (X0 + 9.5, 6.68), (X0 + 9.5, r4 + 0.36)], color=INK, w=1.75)
    _white_label(s, X0 + 6.85, 6.68, 2.8, "no: continue monitoring", size=12)

    # ================================================================== 26  ER diagram
    s = new("Database Design – ER Diagram (PostgreSQL)",
            "Six tables. Machines own alerts and maintenance records; each alert can have one ML feedback row and several maintenance records. Users and ML model metadata are standalone. "
            "Time-series telemetry lives separately in InfluxDB (30 days raw, 1 year aggregated).")
    tw = 2.95
    tx = [X0, X0 + 4.0, X0 + 8.0]
    table_box(s, tx[0], 1.15, tw, "machines", [("id", "PK"), ("name", ""), ("kind", ""), ("ideal_cycle_s", ""), ("temp_limit", ""), ("vib_limit", "")], "blue")
    h_alerts = table_box(s, tx[1], 1.15, tw, "alerts", [("id", "PK"), ("machine_id", "FK"), ("level", ""), ("status", ""), ("opened_at", ""), ("acked_by", ""), ("technician", ""), ("resolution_notes", "")], "orange")
    table_box(s, tx[2], 1.15, tw, "feedback", [("id", "PK"), ("alert_id", "FK"), ("verdict", ""), ("given_by", ""), ("comment", "")], "violet")
    table_box(s, tx[2], 3.2, tw, "users", [("username", "PK"), ("role", ""), ("password_hash", "")], "grey")
    table_box(s, tx[2], 4.7, tw, "ml_models", [("id", "PK"), ("algorithm", ""), ("trained_at", ""), ("threshold", ""), ("metrics", "")], "yellow")
    table_box(s, tx[1], 4.3, tw, "maintenance_logs", [("id", "PK"), ("machine_id", "FK"), ("alert_id", "FK"), ("action", ""), ("technician", ""), ("status", "")], "green")
    line(s, tx[0] + tw, 2.0, tx[1], 2.0, color=INK, w=2, arrow=False)
    tb(s, tx[0] + tw, 1.68, 1.05, 0.26, "1 : N", size=12, bold=True, align="c")
    line(s, tx[1] + tw, 2.0, tx[2], 2.0, color=INK, w=2, arrow=False)
    tb(s, tx[1] + tw, 1.68, 1.05, 0.26, "1 : 1", size=12, bold=True, align="c")
    line(s, tx[1] + tw / 2, 1.15 + h_alerts, tx[1] + tw / 2, 4.3, color=INK, w=2, arrow=False)
    tb(s, tx[1] + tw / 2 + 0.1, 3.7, 0.8, 0.26, "1 : N", size=12, bold=True)
    path_pts = [(tx[0] + tw / 2, 1.15 + 1.9), (tx[0] + tw / 2, 4.9), (tx[1], 4.9)]
    for i in range(len(path_pts) - 1):
        line(s, *path_pts[i], *path_pts[i + 1], color=INK, w=2, arrow=False)
    tb(s, tx[0] + tw / 2 + 0.1, 3.9, 0.8, 0.26, "1 : N", size=12, bold=True)
    box(s, X0, 5.35, 3.4, 1.5, [{"text": "Time series (InfluxDB)", "size": 14, "bold": True}, {"text": "telemetry_raw – 30 days", "size": 13}, {"text": "telemetry_agg – 1 year", "size": 13}],
        fill="cyan", edge="cyan", align="l", margin=0.2)

    # ================================================================== 27  alert lifecycle (state diagram)
    s = new("UML – State Diagram: Alert Lifecycle",
            "An alert is born Open when a detector fires. An operator acknowledges it (or assigns a technician, which acknowledges it) and later resolves it with notes. If the condition disappears "
            "on its own for 3 to 4 seconds, an unacknowledged alert becomes Cleared and an acknowledged one becomes Resolved. ML feedback and maintenance records can be added at any time.")
    cy_ = 2.35
    dotx = X0 + 0.15
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, E(dotx - 0.17), E(cy_ - 0.17), E(0.34), E(0.34))
    o.shadow.inherit = False
    o.fill.solid()
    o.fill.fore_color.rgb = INK
    o.line.fill.background()
    st = [("Open", "status = open", X0 + 0.7, 2.0, "red"), ("Acknowledged", "status = acknowledged", X0 + 4.2, 2.4, "orange"), ("Resolved", "status = resolved", X0 + 8.1, 2.1, "green")]
    for nm, d, x, w, col in st:
        box(s, x, cy_ - 0.55, w, 1.1, [{"text": nm, "size": 18, "bold": True}, {"text": d, "size": 13}], fill=col, edge=col, line=2)
    line(s, dotx + 0.19, cy_, X0 + 0.7, cy_, color=INK, w=2)
    line(s, X0 + 2.7, cy_, X0 + 4.2, cy_, color=INK, w=2)
    tb(s, X0 + 2.7, cy_ - 0.75, 1.5, 0.7, ["acknowledge", "or assign technician"], size=12, align="c", anchor="b")
    line(s, X0 + 6.6, cy_, X0 + 8.1, cy_, color=INK, w=2)
    tb(s, X0 + 6.6, cy_ - 0.75, 1.5, 0.7, ["resolve", "+ notes"], size=12, align="c", anchor="b")
    box(s, X0 + 0.7, 4.1, 2.0, 1.1, [{"text": "Cleared", "size": 18, "bold": True}, {"text": "status = cleared", "size": 13}], fill="grey", edge="grey", line=2)
    line(s, X0 + 1.7, cy_ + 0.55, X0 + 1.7, 4.1, color=INK, w=2)
    tb(s, X0 + 1.85, 3.25, 2.3, 0.6, ["condition ends", "on its own (3 – 4 s)"], size=12, anchor="m")
    path(s, [(X0 + 2.7, cy_ + 0.35), (X0 + 3.45, cy_ + 0.35), (X0 + 3.45, 3.5), (X0 + 9.15, 3.5), (X0 + 9.15, cy_ + 0.55)], color=SLATE, w=1.5, dash=True)
    _white_label(s, X0 + 6.3, 3.5, 4.2, "Open → Resolved is also possible: resolve directly", size=12)
    card(s, X0 + 3.4, 4.1, W - 3.4, 2.2, "grey", "At any state", ["Operator can add **ML feedback** (true / false positive)", "A **maintenance record** can be linked to the alert",
                                                                 "Every action is time-stamped with the user and stored in PostgreSQL"], size=16, title_size=19, gap=8)

    # ================================================================== 28  run modes
    s = new("Flexible Run Modes",
            "The same code runs in four combinations by changing two lines in the .env file. The quick start needs nothing installed except Python and Node: direct ingest and memory storage. "
            "The full path from the SRS uses MQTT through Mosquitto and the two databases.")
    modes = [("INGEST_MODE", "How readings reach the backend", [("direct", "Simulator runs inside the backend – zero infrastructure (quick start, no Docker).", ["Simulator", "Backend"]),
                                                             ("mqtt", "Simulator publishes to Mosquitto over MQTT / WebSockets; backend subscribes (full SRS path).", ["Simulator", "Mosquitto", "Backend"])]),
             ("STORAGE_MODE", "Where data is kept", [("memory  (default)", "Rolling windows and history held in memory; resets on restart. No database needed.", ["Backend", "Memory"]),
                                                  ("db", "InfluxDB for telemetry, PostgreSQL for alerts, maintenance and feedback – survives restarts.", ["Backend", "InfluxDB", "PostgreSQL"])])]
    colx = [X0, X0 + 2.6, X0 + 6.85]
    colw = [2.4, 4.1, 4.1]
    for r, (sw, desc, opts) in enumerate(modes):
        y = 1.2 + r * 2.75
        box(s, colx[0], y, colw[0], 2.5, [{"text": sw, "size": 15, "bold": True, "space_after": 6}, {"text": desc, "size": 14}], fill="blue", edge="blue", align="l", margin=0.2)
        for c, (nm, d, flow) in enumerate(opts):
            x = colx[c + 1]
            box(s, x, y, colw[c + 1], 2.5, "", fill="grey", edge="white")
            tb(s, x + 0.2, y + 0.12, colw[c + 1] - 0.4, 0.35, nm, size=17, bold=True)
            tb(s, x + 0.2, y + 0.55, colw[c + 1] - 0.4, 0.8, d, size=13)
            n = len(flow)
            fw = 1.05 if n == 3 else 1.3
            gap_ = 0.3
            total = n * fw + (n - 1) * gap_
            fx = x + (colw[c + 1] - total) / 2
            for k, f in enumerate(flow):
                colr = ["blue", "orange", "green", "cyan", "violet"][k % 5]
                box(s, fx + k * (fw + gap_), y + 1.6, fw, 0.65, f, fill=colr, edge=colr, size=12, bold=True, margin=0.03)
                if k < n - 1:
                    line(s, fx + k * (fw + gap_) + fw + 0.02, y + 1.925, fx + (k + 1) * (fw + gap_) - 0.02, y + 1.925, color=INK, w=1.75)
    tb(s, X0, 6.7, W, 0.3, "Same code, same dashboard – only two lines in .env change.", size=14, color=GREEN_D, bold=True)

    # ================================================================== 29  docker + requirements
    s = new("Deployment – Docker Compose & Requirements (SRS §10–11)",
            "One command starts the stack: Mosquitto, simulator, backend and the nginx-served React frontend. The two databases are an optional compose profile. "
            "The databases were verified live in Docker containers. The right panel is the hardware and software requirement table from the SRS.")
    box(s, X0, 1.1, 6.6, 0.55, "docker compose up --build", fill="grey", edge="grey", size=18, bold=True, color=INK)
    sv = cols(3, gap=0.375, x0=X0, w=6.6)
    names3 = [("simulator", "publishes telemetry", "blue"), ("mosquitto", ":1883 · :9001", "orange"), ("backend", ":8000 · /docs", "green")]
    for (x, w), (nm, d, col) in zip(sv, names3):
        box(s, x, 1.9, w, 1.1, [{"text": nm, "size": 15, "bold": True}, {"text": d, "size": 12}], fill=col, edge=col, margin=0.05)
    line(s, sv[0][0] + sv[0][1] + 0.01, 2.45, sv[1][0] - 0.01, 2.45, color=INK, w=2)
    line(s, sv[1][0] + sv[1][1] + 0.01, 2.45, sv[2][0] - 0.01, 2.45, color=INK, w=2)
    box(s, sv[2][0], 3.75, sv[2][1], 1.1, [{"text": "frontend", "size": 15, "bold": True}, {"text": "nginx · :5173", "size": 12}], fill="cyan", edge="cyan", margin=0.05)
    line(s, sv[2][0] + 1.0, 3.74, sv[2][0] + 1.0, 3.01, color=INK, w=2, back_arrow=True)
    box(s, X0, 3.75, 4.35, 1.1, "", fill=None, edge="grey", dash=True, radius=0.1)
    tb(s, X0 + 0.15, 3.8, 4.1, 0.3, "optional profile  --profile db", size=12, color=MUTED, italic=True)
    for k, (nm, d) in enumerate([("PostgreSQL 14", ":5432"), ("InfluxDB 2", ":8086")]):
        box(s, X0 + 0.15 + k * 2.1, 4.15, 1.95, 0.6, [{"text": nm, "size": 13, "bold": True}, {"text": d, "size": 11}], fill="violet", edge="violet", margin=0.03)
    path(s, [(sv[2][0] + 0.3, 3.01), (sv[2][0] + 0.3, 3.4), (X0 + 2.18, 3.4), (X0 + 2.18, 3.74)], color=INK, w=2)
    card(s, X0, 5.15, 6.6, 1.7, "green", "Verified live", ["PostgreSQL 14 and InfluxDB 2 running in Docker: telemetry, alerts, acknowledgements, feedback and maintenance saved and still there after a backend restart"], size=14, bullets=False)
    rxx, rww = X0 + 6.85, W - 6.85
    box(s, rxx, 1.1, rww, 5.75, "", fill="grey", edge="white")
    tb(s, rxx + 0.2, 1.22, rww - 0.4, 0.4, "Requirements (SRS §11)", size=17, bold=True)
    tb(s, rxx + 0.2, 1.75, rww - 0.4, 2.0, [{"text": "Hardware – minimum → recommended", "size": 14, "bold": True, "space_after": 4},
                                           {"text": "CPU: dual-core 2.0 GHz → quad-core 2.5 GHz", "size": 13, "bullet": True, "space_after": 3},
                                           {"text": "RAM: 8 GB → 16 GB", "size": 13, "bullet": True, "space_after": 3},
                                           {"text": "Storage: 20 GB free → 50 GB SSD", "size": 13, "bullet": True, "space_after": 3},
                                           {"text": "Display: 1366 × 768 → 1920 × 1080", "size": 13, "bullet": True}], size=13)
    tb(s, rxx + 0.2, 3.85, rww - 0.4, 2.9, [{"text": "Software", "size": 14, "bold": True, "space_after": 4},
                                           {"text": "Windows 10/11, Ubuntu 20.04+, macOS 11+", "size": 13, "bullet": True, "space_after": 3},
                                           {"text": "Docker Desktop 4.x + Compose 2.x", "size": 13, "bullet": True, "space_after": 3},
                                           {"text": "Python 3.10+ and Node 18+ (containerised)", "size": 13, "bullet": True, "space_after": 3},
                                           {"text": "PostgreSQL 14, InfluxDB 2, Mosquitto 2", "size": 13, "bullet": True, "space_after": 3},
                                           {"text": "Chrome, Firefox or Edge", "size": 13, "bullet": True}], size=13)
